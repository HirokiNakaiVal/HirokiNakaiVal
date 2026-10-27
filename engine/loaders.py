"""取り込みファイルの種類判定と読み込み（正規化したDataFrameを返す）"""
import re, datetime as dt
import pandas as pd
import pdfplumber
import openpyxl
from .common import WD

KINDS = {
    "plan": "月間稼働表",
    "taskrelay": "タスクリレー",
    "kadou": "KADOU実績",
    "sales_hour": "時間帯別売上",
    "sales_reg": "時間帯別レジ別売上",
    "budget": "日割予算",
    "labor_budget": "労働時間予算",
}

class LoadError(Exception):
    pass

# ---------- 種類判定 ----------
KADOU_COLS = {"shain_nm", "ymd", "paint_sbt2_nm", "s_time"}

def _csv_header(path):
    for enc in ("utf-8-sig", "cp932"):
        try:
            with open(path, encoding=enc) as f:
                return enc, [c.strip().strip('"') for c in f.readline().split(",")]
        except UnicodeDecodeError:
            continue
    raise LoadError("CSVの文字コードが判別できません")

def classify(path):
    low = path.lower()
    if low.endswith(".csv"):
        _, hdr = _csv_header(path)
        if KADOU_COLS <= set(hdr): return "kadou"
        raise LoadError("種類を判別できないCSVです（KADOU実績の列：shain_nm・ymd・paint_sbt2_nm・s_time）")
    if low.endswith(".pdf"):
        with pdfplumber.open(path) as pdf:
            t = (pdf.pages[0].extract_text() or "")
        if "月間稼働表" in t: return "plan"
        if "日割予算" in t: return "budget"
        if "日別稼働表" in t: raise LoadError("日別稼働表PDFは取り込み対象外です（月間稼働表を取り込んでください）")
        raise LoadError("種類を判別できないPDFです")
    if low.endswith(".xls"):
        if is_labor_budget(path): return "labor_budget"
        raise LoadError("種類を判別できないExcel(.xls)です")
    if low.endswith((".xlsx", ".xlsm")):
        if is_labor_budget(path): return "labor_budget"
        wb = openpyxl.load_workbook(path, read_only=True)
        if "勤務データ" in wb.sheetnames: return "plan"
        ws = wb.worksheets[0]
        rows = list(ws.iter_rows(min_row=1, max_row=2, values_only=True))
        a1 = str(rows[0][0] or "") if rows else ""
        hdr = [str(x) for x in (rows[1] if len(rows) > 1 else []) if x is not None]
        if "KADOU" in a1 or {"shain_nm", "paint_sbt2_nm", "s_time"} <= set(hdr): return "kadou"
        if "TaskRelay" in a1 or {"作業指示ID", "作業者氏名"} <= set(hdr): return "taskrelay"
        if {"取引日", "時間帯", "客数"} <= set(hdr):
            return "sales_reg" if "レジNo" in hdr else "sales_hour"
        raise LoadError("種類を判別できないExcelです")
    raise LoadError("PDF・Excel(.xlsx/.xls)・CSVを取り込んでください")

# ---------- 月間稼働表PDF ----------
TIME = re.compile(r"^\d{2}:\d{2}$")
WORDS = ["公休", "有休", "会議", "休職", "研修", "特休", "慶弔", "欠勤", "出産", "育児", "介護", "労災", "希望休"]

def parse_plan_pdf(path):
    rows, events = [], {}
    with pdfplumber.open(path) as pdf:
        head = (pdf.pages[0].extract_text() or "").splitlines()[0]
        mm = re.search(r"(\d{4}/\d{2}/\d{2})\s*-\s*(\d{4}/\d{2}/\d{2})", head)
        if not mm: raise LoadError("月間稼働表の期間が読み取れません")
        d0, d1 = pd.Timestamp(mm.group(1)), pd.Timestamp(mm.group(2))
        store = (re.search(r"所属：(\S+)", head) or [None, ""])[1]
        company = (re.search(r"会社：(\S+)", head) or [None, ""])[1]
        days = list(pd.date_range(d0, d1))
        dept = None
        for page in pdf.pages:
            words = page.extract_words()
            ys = {}
            for w in words:
                if re.fullmatch(r"\d{2}", w["text"]): ys.setdefault(round(w["top"]), []).append(w)
            if not ys: continue
            drow = max(ys.values(), key=len)
            if len(drow) != len(days): raise LoadError(f"日付列の数が期間と合いません（{len(drow)}列／{len(days)}日）")
            drow.sort(key=lambda w: w["x0"])
            cx = [(w["x0"] + w["x1"]) / 2 for w in drow]
            half = (cx[1] - cx[0]) / 2
            date_top = drow[0]["top"]
            def col(x):
                i = min(range(len(cx)), key=lambda k: abs(cx[k] - x))
                return i if abs(cx[i] - x) <= half else None
            if page.page_number == 1:      # 行事（縦書き）
                chars = [c for c in page.chars if 50 < c["top"] < date_top - 5 and c["x0"] > cx[0] - half]
                for i in range(len(days)):
                    subs = {}
                    for c in chars:
                        if col((c["x0"] + c["x1"]) / 2) == i: subs.setdefault(round(c["x0"] / 6), []).append(c)
                    t = " ".join("".join(c["text"] for c in sorted(subs[k], key=lambda c: c["top"])).strip()
                                 for k in sorted(subs)).strip()
                    if t: events[days[i]] = t
            # 部門見出し＝横幅いっぱいの背景矩形がある行
            bands = [(r["top"], r["bottom"]) for r in page.rects if r["width"] > (cx[-1] - cx[0])]
            halfs = [c for c in page.curves if c.get("non_stroking_color") and len(c.get("pts", [])) >= 4
                     and len(c["non_stroking_color"]) == 3 and abs(c["non_stroking_color"][2] - 0.835) < 0.05]
            left = {}
            for w in words:
                if w["x1"] < cx[0] - half and w["top"] > date_top + 20:
                    left.setdefault(round(w["top"]), []).append(w)
            for y in sorted(left):
                txt = " ".join(w["text"] for w in sorted(left[y], key=lambda w: w["x0"])).replace("　", " ")
                if any(k in txt for k in ("計画人数", "社員名", "行事", "Powered", "合計", "ページ")) or txt == "日付":
                    continue
                if any(a - 2 <= y <= b for a, b in bands):
                    dept = txt; continue
                cell = {}
                for w in words:
                    if abs(w["top"] - y) <= 14 and w["x0"] >= cx[0] - half:
                        i = col((w["x0"] + w["x1"]) / 2)
                        if i is not None: cell.setdefault(i, []).append(w)
                hset = set()
                for c in halfs:
                    if abs((c["top"] + c["bottom"]) / 2 - y) <= 16:
                        i = col((c["x0"] + c["x1"]) / 2)
                        if i is not None: hset.add(i)
                for i, d in enumerate(days):
                    toks = [w["text"] for w in sorted(cell.get(i, []), key=lambda w: w["top"])]
                    times = [t for t in toks if TIME.match(t)]
                    word = next((t for t in toks if t in WORDS), None)
                    st = en = None
                    if len(times) >= 2:
                        kubun = "半休" if i in hset else "出勤"
                        st = dt.time(*map(int, times[0].split(":"))); en = dt.time(*map(int, times[1].split(":")))
                    elif word: kubun = word
                    else: kubun = "未入力" if not toks else toks[0]
                    rows.append(dict(日付=d, 曜日=WD[d.weekday()], 部門=dept, 社員名=txt, 区分=kubun,
                                     開始=st, 終了=en, 原表記=" ".join(toks) or None))
    df = pd.DataFrame(rows)
    if df.empty: raise LoadError("社員の行が読み取れませんでした")
    # 同名で全日公休の重複行（兼務など）を除く
    drop = []
    for nm, x in df.groupby("社員名"):
        if len(x) > x["日付"].nunique():
            x = x.assign(n=x.groupby("日付").cumcount())
            for _, y in x.groupby("n"):
                if (y["区分"] == "公休").all(): drop += list(y.index)
    df = df.drop(drop).reset_index(drop=True)
    df["行事"] = df["日付"].map(events)
    info = dict(store=store, company=company, period_from=d0, period_to=d1,
                employees=int(df["社員名"].nunique()), removed_dup_rows=len(drop))
    return df, info

def parse_plan_excel(path):
    df = pd.read_excel(path, sheet_name="勤務データ")
    need = ["日付", "部門", "社員名", "区分", "開始", "終了"]
    if not set(need) <= set(df.columns): raise LoadError("勤務データシートの列が足りません")
    if "行事" not in df: df["行事"] = None
    df["日付"] = pd.to_datetime(df["日付"])
    for c in ("開始", "終了"):
        df[c] = df[c].apply(lambda v: v if isinstance(v, dt.time) else None)
    return df[need + ["行事"]], dict(period_from=df["日付"].min(), period_to=df["日付"].max(),
                                     employees=int(df["社員名"].nunique()))

# ---------- Excel系 ----------
def _nm(s): return s.astype(str).str.replace("　", " ").str.strip()

def parse_kadou(path):
    if path.lower().endswith(".csv"):
        enc, _ = _csv_header(path)
        k = pd.read_csv(path, encoding=enc, dtype=str)
    else:
        k = pd.read_excel(path, header=1, dtype=str)
    rows_in = len(k)
    # 抽出によっては同じ行が何十回も重複している（DB実績CSVで約19倍）→ 完全に同じ行は1行に
    k = k.drop_duplicates()
    dup = rows_in - len(k)
    k["s_time"] = pd.to_numeric(k["s_time"], errors="coerce")
    shozoku = sorted(k["shozoku_cd"].dropna().astype(str).str.strip().unique().tolist()) if "shozoku_cd" in k else []
    k = k.rename(columns={"paint_sbt2_nm": "item"})
    k["nm"] = _nm(k["shain_nm"])
    k["ymd"] = pd.to_datetime(k["ymd"])
    k["item"] = k["item"].astype(str).str.strip()
    k = k[["nm", "ymd", "item", "s_time"]].dropna()
    # 「小林 久範(登販)」のように資格などの括弧書き付きで同じ人が別IDで登録され、同じ日の同じ作業が重複して
    # 出ることがある → 括弧を外した名前が同じで同じ日なら、括弧の無い方（無ければ最初の名前）だけを残す
    base = k["nm"].str.replace(r"\s*[（(][^（）()]*[）)]\s*$", "", regex=True).str.strip()
    k = k.assign(_base=base)
    keep = {}
    for (b, d), names in k.groupby(["_base", "ymd"])["nm"].unique().items():
        keep[(b, d)] = b if b in names else names[0]
    sel = [keep[(b, d)] == n for b, d, n in zip(k["_base"], k["ymd"], k["nm"])]
    alias_rows = int(len(k) - sum(sel))
    k = k[sel].copy()
    k["nm"] = k["_base"]
    k = k.drop(columns="_base")
    # 同じ作業名が課ごとに分かれていることがある（1課の接客・2課の接客）→ 社員・日付・作業で合計
    k = k.groupby(["nm", "ymd", "item"], as_index=False)["s_time"].sum()
    return k, dict(period_from=k["ymd"].min(), period_to=k["ymd"].max(), employees=int(k["nm"].nunique()),
                   rows=rows_in, removed_dup_rows=dup, removed_alias_rows=alias_rows, shozoku=",".join(shozoku[:3]))

def parse_taskrelay(path):
    t = pd.read_excel(path, header=1)
    t["nm"] = _nm(t["作業者氏名"])
    for c in ("実施期間F", "実施期間TO"): t[c] = pd.to_datetime(t[c])
    keep = ["作業指示ID", "作業者氏名", "nm", "タイトル", "重要度", "実施期間F", "実施期間TO", "作業予定時間",
            "作業時間実績", "発信者組織名", "指示割当作成者名"]
    t = t[[c for c in keep if c in t.columns]]
    return t, dict(period_from=t["実施期間F"].min(), period_to=t["実施期間TO"].max(), rows=len(t))

def parse_sales(path, kind):
    s = pd.read_excel(path, header=1)
    s["取引日"] = pd.to_datetime(s["取引日"]); s["h"] = s["時間帯"].astype(int)
    cols = ["取引日", "h", "売上高", "客数"] + (["レジNo"] if kind == "sales_reg" else [])
    s = s[cols]
    return s, dict(period_from=s["取引日"].min(), period_to=s["取引日"].max(), rows=len(s))

# ---------- 労働時間予算（店番×月、正社員・パート・アルバイト）----------
MONTH_COLS = ["3月", "4月", "5月", "6月", "7月", "8月", "9月", "10月", "11月", "12月", "1月", "2月"]
_Z2H = str.maketrans("０１２３４５６７８９", "0123456789")

def _raw_sheet(path):
    return pd.read_excel(path, header=None, sheet_name=0, dtype=object)

def is_labor_budget(path):
    try:
        df = _raw_sheet(path)
    except Exception:
        return False
    head = " ".join(str(v) for v in df.head(6).values.ravel() if isinstance(v, str))
    return "労働時間予算" in head

def parse_labor_budget(path):
    """戻り値：店番・月(期の何か月目 0〜11)・正社員・パート・アルバイト。期は『１２期』などの見出しから"""
    df = _raw_sheet(path)
    title = next((str(v) for v in df.head(6).values.ravel() if isinstance(v, str) and "労働時間予算" in v), "")
    ki = (re.search(r"([0-9０-９]+)\s*期", title) or [None, None])[1]
    rows, cur, mcol = [], None, None
    for _, r in df.iterrows():
        v0, v1 = r.iloc[0], r.iloc[1]
        if isinstance(v0, str) and v0.strip() == "店番":
            mcol = {}
            for j, h in enumerate(r.iloc[2:], start=2):
                h = str(h).translate(_Z2H).strip()
                if h in MONTH_COLS: mcol[MONTH_COLS.index(h)] = j
            continue
        if isinstance(v0, (int, float)) and not pd.isna(v0): cur = int(v0)
        if cur is None or mcol is None or not isinstance(v1, str): continue
        kind = {"正社員予算": "正社員", "パート予算": "パート", "アルバイト予算": "アルバイト"}.get(v1.strip())
        if not kind: continue
        for mi, j in mcol.items():
            x = r.iloc[j]
            x = float(x) if isinstance(x, (int, float)) and not pd.isna(x) else 0.0
            rows.append(dict(店番=cur, 月=mi, 区分=kind, 時間=x))
    if not rows: raise LoadError("労働時間予算の表が読み取れません")
    t = pd.DataFrame(rows).pivot_table(index=["店番", "月"], columns="区分", values="時間", aggfunc="sum").fillna(0).reset_index()
    for k in ("正社員", "パート", "アルバイト"):
        if k not in t: t[k] = 0.0
    t["合計"] = t["正社員"] + t["パート"] + t["アルバイト"]
    return t, dict(ki=ki, stores=int(t["店番"].nunique()), title=title.strip())

# ---------- 日割予算PDF ----------
def parse_budget(path):
    with pdfplumber.open(path) as pdf:
        pages = [(p.extract_text() or "") for p in pdf.pages]
    txt = "\n".join(pages)
    ym = re.search(r"(20\d\d)\s*年\s*(\d{1,2})\s*月分", txt)
    if not ym: raise LoadError("日割予算の年月が読み取れません")
    y, mo = int(ym.group(1)), int(ym.group(2))
    out = {}
    for pt in pages:
        if not re.search(r"合\s*計\s+客数", pt): continue        # 合計列のあるページ
        for line in pt.splitlines():
            dm = re.match(r"\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日\s*\S?\s+(.*)$", line)
            if not dm or int(dm.group(1)) != mo: continue
            nums = re.findall(r"-?[\d,]+", dm.group(3))
            if len(nums) >= 3: out[pd.Timestamp(y, mo, int(dm.group(2)))] = int(nums[-3].replace(",", "")) * 1000
    if not out: raise LoadError("日割予算の合計列が読み取れません")
    g = lambda pat: (re.search(pat, txt) or [None, ""])[1].strip() if re.search(pat, txt) else ""
    df = pd.DataFrame(sorted(out.items()), columns=["日付", "予算売上"])
    return df, dict(period_from=df["日付"].min(), period_to=df["日付"].max(), days=len(df),
                    company=g(r"会社\s*：\s*\d+\s*(\S+)"), division=g(r"事業部\s*：\s*\d+\s*(\S+)"),
                    store=g(r"店舗\s*：\s*(\d+\s*\S+)"))

def load(path, kind=None):
    kind = kind or classify(path)
    if kind == "plan":
        df, info = parse_plan_pdf(path) if path.lower().endswith(".pdf") else parse_plan_excel(path)
    elif kind == "kadou": df, info = parse_kadou(path)
    elif kind == "taskrelay": df, info = parse_taskrelay(path)
    elif kind in ("sales_hour", "sales_reg"): df, info = parse_sales(path, kind)
    elif kind == "budget": df, info = parse_budget(path)
    elif kind == "labor_budget": df, info = parse_labor_budget(path)
    else: raise LoadError(f"未対応の種類: {kind}")
    return kind, df, info

# ---------- 店舗データの結合 ----------
class StoreData:
    """同じ種類のファイルが複数あれば結合（新しく取り込んだものを優先）"""
    def __init__(self, frames):          # frames: [(kind, df)] 取込順（古い→新しい）
        by = {}
        for kind, df in frames:
            if kind == "plan":
                # 月間稼働表の行順（部署順・社員順）を「順」として残す。ファイルは社員ごとに日付が並ぶので
                # 社員が最初に出てくる位置がそのまま表示順になる
                df = df.copy()
                first = {}
                for i, nm in enumerate(df["社員名"]):
                    first.setdefault(nm, i)
                df["順"] = df["社員名"].map(first)
            by.setdefault(kind, []).append(df)
        def cat(kind, keys):
            if kind not in by: return None
            d = pd.concat(by[kind], ignore_index=True)
            return d.drop_duplicates(keys, keep="last").reset_index(drop=True)
        self.plan = cat("plan", ["日付", "社員名"])
        self.kadou = None
        if "kadou" in by:
            parts = [d.assign(_f=i) for i, d in enumerate(by["kadou"])]
            k = pd.concat(parts, ignore_index=True)
            last = k.groupby(["nm", "ymd"])["_f"].transform("max")      # 同じ社員・同じ日は最後に取り込んだファイルだけ
            self.kadou = k[k["_f"] == last].drop(columns="_f").reset_index(drop=True)
        self.tasks = cat("taskrelay", ["作業指示ID"])
        self.sales_hour = cat("sales_hour", ["取引日", "h"])
        self.sales_reg = cat("sales_reg", ["取引日", "h", "レジNo"])
        self.budget = cat("budget", ["日付"])
