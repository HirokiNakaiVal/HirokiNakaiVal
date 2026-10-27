"""DayPlanner の結果を Excel（日別稼働表・レジ別稼働表ほか）に書き出す"""
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment
from .common import SLOT, WD, hm, isreg

FONT = "Meiryo UI"
G0, G1 = 6 * 60, 22 * 60
NS = (G1 - G0) // SLOT
GC = 4
thin = Side(style="thin", color="A6A6A6")
hair = Side(style="hair", color="D0D0D0")
C = Alignment(horizontal="center", vertical="center", wrap_text=False, shrink_to_fit=True)
def F(**k): return Font(name=FONT, size=k.pop("size", 9), **k)
PCOL = ["FFE699", "C6E0B4", "BDD7EE", "F8CBAD", "D9C3E9", "B4DEE6", "FFD966", "A9D08E",
        "9BC2E6", "F4B084", "C9C9C9", "E2EFDA", "FCE4D6", "DDEBF7", "FFF2CC", "EDEDED"]

def color(a, kero):
    if a.startswith("TR:"): return "F2C4DE"
    if a == "休憩": return "D9D9D9"
    if a == "接客": return "DDEBF7"
    if a == "タイヤ": return "C6E0B4"
    if kero and a == kero: return "F8CBAD"
    if a in ("朝礼", "昼礼"): return "D9C3E9"
    if isreg(a) or a in ("精算", "釣銭準備"): return "FFE699"
    if a.startswith("SC"): return "B4DEE6"
    return "E2EFDA"

def blocks(seq):
    out = []
    for T in sorted(seq):
        if out and out[-1][1] == T and out[-1][2] == seq[T]: out[-1][1] = T + 1
        else: out.append([T, T + 1, seq[T]])
    return out

def write(pl, path):
    md = f"{pl.day.month}/{pl.day.day}"; wdn = WD[pl.day.weekday()]
    store = pl.store_name
    tr = pl.tr
    title_of = dict(zip(tr["作業指示ID"].astype(int), tr["タイトル"])) if len(tr) else {}
    wb = Workbook()
    wp = wb.active; wp.title = "日別稼働表"
    wr = wb.create_sheet("レジ別稼働表")
    wf = wb.create_sheet("時間帯予測")
    wt = wb.create_sheet("タスクリレー割当")
    wl = wb.create_sheet("レジ配置")
    wa = wb.create_sheet("作業配置")
    we = wb.create_sheet("作業推定")
    wn = wb.create_sheet("説明")

    def grid_header(w, r, title):
        w.cell(r, 1, title).font = F(bold=True, color="FFFFFF")
        for c in range(1, GC + NS): w.cell(r, c).fill = PatternFill("solid", fgColor="305496")
        for h in range(G0 // 60, G1 // 60):
            c = GC + (h * 60 - G0) // SLOT
            w.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + 3)
            x = w.cell(r, c, f"{h}時"); x.font = F(bold=True, color="FFFFFF"); x.alignment = C
    def row_border(w, r):
        for c in range(1, GC + NS):
            w.cell(r, c).border = Border(top=thin, bottom=thin, left=thin if c < GC or (c - GC) % 4 == 0 else hair)
    def paint(w, r, T0, T1, label, fill, comment=None):
        T0 = max(T0, G0 // SLOT); T1 = min(T1, G1 // SLOT)
        if T1 <= T0: return
        c0, c1 = GC + T0 - G0 // SLOT, GC + T1 - 1 - G0 // SLOT
        for c in range(c0, c1 + 1): w.cell(r, c).fill = PatternFill("solid", fgColor=fill)
        if c1 > c0: w.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c1)
        x = w.cell(r, c0, label); x.font = F(size=8); x.alignment = C
        if comment: x.comment = Comment(comment, "稼働")
    def hourband(w, r, label, vals, fmt="0"):
        w.cell(r, 1, label).font = F(bold=True)
        for h in range(G0 // 60, G1 // 60):
            c = GC + (h * 60 - G0) // SLOT
            w.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + 3)
            x = w.cell(r, c, vals(h)); x.font = F(size=8); x.alignment = C; x.number_format = fmt
    for w in (wr, wp):
        w.column_dimensions["A"].width = 14; w.column_dimensions["B"].width = 6; w.column_dimensions["C"].width = 6
        for i in range(NS): w.column_dimensions[get_column_letter(GC + i)].width = 2.6

    # ---- 作業配置（他シートの数式の元）
    wa.append(["部門", "社員名", "作業", "開始", "終了", "分", "休憩の根拠", "指示ID", "タスクリレー指示"])
    i = 2
    for nm, p in pl.people.items():
        if p["st"] is None: continue
        for T0, T1, a in blocks(p["slots"]):
            lab, tid = a, None
            if a.startswith("TR:"): lab, tid = "タスクリレー", int(a[3:])
            wa.append([p["dept"], nm, lab, dt.time(T0 * SLOT // 60, T0 * SLOT % 60),
                       dt.time(T1 * SLOT // 60 % 24, T1 * SLOT % 60),
                       f"=ROUND((E{i}-D{i})*1440,0)", p["bsrc"] if a == "休憩" else None, tid, title_of.get(tid)])
            wa.cell(i, 4).number_format = wa.cell(i, 5).number_format = "h:mm"; i += 1
    NA = max(i - 1, 2)

    # ---- レジ配置
    lrows = []
    for rg in pl.REGS:
        for T0, T1, nm in blocks(pl.occ[rg]): lrows.append([rg, nm, pl.people[nm]["dept"], T0 * SLOT, T1 * SLOT])
    posts = [x for x in (pl.kero, "SC1", "釣銭準備", "精算") if x]
    for nm, p in pl.people.items():
        if p["st"] is None: continue
        for T0, T1, a in blocks(p["slots"]):
            if a in posts: lrows.append([a, nm, p["dept"], T0 * SLOT, T1 * SLOT])
    wl.append(["レジ・持ち場", "社員名", "部門", "開始", "終了", "分"])
    for j, (rg, nm, d, s_, e_) in enumerate(lrows, start=2):
        wl.append([rg, nm, d, dt.time(s_ // 60, s_ % 60), dt.time(e_ // 60 % 24, e_ % 60), f"=ROUND((E{j}-D{j})*1440,0)"])
        wl.cell(j, 4).number_format = wl.cell(j, 5).number_format = "h:mm"
    NL = max(len(lrows) + 1, 2)
    LR = lambda col: f"レジ配置!${col}$2:${col}${NL}"

    # ---- 時間帯予測
    hours = list(range(pl.OPEN // 60, pl.CLOSE // 60 + 1))
    wf.append(["時間帯", "基準日平均 客数", "予測 客数", "予測 有人レジ客数", "1台あたり処理客数(目安)", "必要レジ台数",
               "配置レジ台数(平均)", "1台あたり予測客数", "判定", f"{md}実績 客数(参考)"])
    FR = {}
    last = f"{pl.CLOSE // 60}時"
    for j, h in enumerate(hours, start=2):
        FR[h] = j
        hs_, he_ = h * 60 / 1440, (h + 1) * 60 / 1440
        A, D, E = LR("A"), LR("D"), LR("E")
        ov = (f"SUMPRODUCT((RIGHT({A},1)=\"R\")*({D}<{he_})*({E}>{hs_})"
              f"*((({E}<{he_})*{E}+({E}>={he_})*{he_})-(({D}>{hs_})*{D}+({D}<={hs_})*{hs_})))*24")
        base_staffed = float(pl.fc_staffed.get(h, 0)) / pl.SCALE if pl.SCALE else 0
        act = pl.act_h.get(h) if len(pl.act_h) else None
        wf.append([f"{h}時", round(float(pl.prof_h.get(h, 0)), 1), f"=B{j}*$M$3", f"={round(base_staffed, 2)}*$M$3", "=$M$4",
                   f'=IF(A{j}="{last}",0,MAX($M$5,ROUNDUP(D{j}/E{j},0)))', f"={ov}",
                   f'=IF(G{j}>0,D{j}/G{j},"")', f'=IF(A{j}="{last}","",IF(G{j}+0.25<F{j},"不足",IF(G{j}>F{j}+1.5,"余裕","")))',
                   None if act is None else int(act)])
        for c in (3, 4, 7, 8): wf.cell(j, c).number_format = "0.0"
        wf.cell(j, 2).font = Font(color="0000FF")
    n_f = len(hours) + 1
    wf.append(["合計", f"=SUM(B2:B{n_f})", f"=SUM(C2:C{n_f})", f"=SUM(D2:D{n_f})", "", "", f"=SUM(G2:G{n_f})", "", "",
               f"=SUM(J2:J{n_f})"])
    for c in (2, 3, 4, 7): wf.cell(n_f + 1, c).number_format = "#,##0"
    wf.append([]); wf.append(["レジ別 予測客数（人/時）"])
    cols = list(pl.fc_r.columns)
    wf.append(["時間帯"] + [f"レジ{n}" + ("" if n in pl.staffed else "(無人)") for n in cols])
    for h in hours:
        wf.append([f"{h}時"] + [round(float(pl.fc_r.loc[h, n]), 1) if h in pl.fc_r.index else 0 for n in cols])
    # 前提値（表の後に書く）
    wf["L1"] = "予算売上(円)"; wf["M1"] = pl.BUD if pl.BUD else (round(pl.ref_sales) if pl.has_sales else 0)
    wf["L2"] = "基準日の平均売上(円)"; wf["M2"] = round(pl.ref_sales) if pl.has_sales else 0
    wf["L3"] = "補正率"; wf["M3"] = "=IF(M2>0,M1/M2,1)"
    wf["L4"] = "1台あたり処理客数(人/時)"; wf["M4"] = float(pl.S["cap"])
    wf["L5"] = "最低レジ台数"; wf["M5"] = int(pl.S["min_lanes"])
    wf["L6"] = "基準日"; wf["M6"] = "、".join(f"{x:%m/%d}" for x in pl.REF) or "（売上実績なし）"
    # 労働時間予算との比較（計画実働は作業配置から数式で集計）
    wf["L8"] = (pl.labor["getsudo"] + "の労働時間予算（h）") if pl.labor else "労働時間予算（月度・h）"; wf["M8"] = round(pl.labor["month"], 1) if pl.labor else None
    wf["L9"] = "日割の比率"; wf["M9"] = round(pl.labor["ratio"], 5) if pl.labor else None
    wf["L10"] = "労働時間予算（日割・h）"; wf["M10"] = "=IF(M8=\"\",\"\",M8*M9)" if pl.labor else None
    wf["L11"] = "計画実働（h）"
    wf["M11"] = f'=(SUM(作業配置!$F$2:$F${NA})-SUMIFS(作業配置!$F$2:$F${NA},作業配置!$C$2:$C${NA},"休憩"))/60'
    wf["L12"] = "差（計画−予算・h）"; wf["M12"] = '=IF(M10="","",M11-M10)'
    wf["L13"] = "日割の方法"; wf["M13"] = pl.labor["method"] if pl.labor else "労働時間予算なし"
    for c in ("M10", "M11", "M12"): wf[c].number_format = "0.0"
    for c in ("M1", "M2", "M4", "M5", "M8", "M9"): wf[c].font = Font(color="0000FF")
    wf["M1"].number_format = wf["M2"].number_format = "#,##0"; wf["M3"].number_format = "0.000"
    ref = lambda col, h: f"=時間帯予測!{col}{FR[h]}" if h in FR else None

    # ---- 日別稼働表（社員別）
    bud = f"　予算 {pl.BUD/1000:,.0f}千円" if pl.BUD else ""
    lab = (f"　労働時間 予算{pl.labor['day']:.1f}h／計画{pl.planned_hours:.1f}h" if pl.labor else f"　計画実働{pl.planned_hours:.1f}h")
    wp["A1"] = f"日別稼働表　{store}　{pl.day:%Y/%m/%d}({wdn}){bud}{lab}" + (f"　行事：{pl.event_today}" if pl.event_today else "")
    wp["A1"].font = F(size=12, bold=True)
    r = 3
    hourband(wp, r, "予測客数", lambda h: f"=ROUND(時間帯予測!C{FR[h]},0)" if h in FR else None); r += 1
    for dept, g in pl.plan.groupby("部門", sort=False):
        grid_header(wp, r, f"{pl.day:%d}({wdn}) {dept}"); r += 1
        for nm in g["社員名"]:
            p = pl.people[nm]; wp.cell(r, 1, nm).font = F(); row_border(wp, r)
            if p["st"] is None:
                wp.cell(r, 2, p["kubun"]).font = F(color="7F7F7F")
            else:
                wp.cell(r, 2, hm(p["st"])).font = F(); wp.cell(r, 3, hm(p["en"])).font = F()
                if p["kubun"] == "半休": wp.cell(r, 1).comment = Comment("半休（月間稼働表）", "稼働")
                for T0, T1, a in blocks(p["slots"]):
                    lab, cm = a, None
                    if a.startswith("TR:"): lab, cm = "タスクリレー", f"{a[3:]}\n{title_of.get(int(a[3:]), '')}"
                    paint(wp, r, T0, T1, lab, color(a, pl.kero), cm)
            r += 1
        r += 1
    wp.freeze_panes = "D4"

    # ---- レジ別稼働表
    wr["A1"] = f"レジ別稼働表　{store}　{pl.day:%Y/%m/%d}({wdn})　営業 {hm(pl.OPEN)}〜{hm(pl.CLOSE)}"
    wr["A1"].font = F(size=12, bold=True)
    r = 3
    hourband(wr, r, "予測客数", lambda h: f"=ROUND(時間帯予測!C{FR[h]},0)" if h in FR else None); r += 1
    hourband(wr, r, "必要レジ台数", lambda h: ref("F", h)); r += 1
    hourband(wr, r, "配置レジ台数", lambda h: f"=ROUND(時間帯予測!G{FR[h]},1)" if h in FR else None, "0.0"); r += 1
    hourband(wr, r, "判定", lambda h: ref("I", h)); r += 2
    grid_header(wr, r, f"{pl.day:%d}({wdn}) レジ"); wr.cell(r, 2, "稼働h").font = F(bold=True, color="FFFFFF"); r += 1
    pcol = {}
    def colp(nm):
        if nm not in pcol: pcol[nm] = PCOL[len(pcol) % len(PCOL)]
        return pcol[nm]
    short = lambda nm: nm.split(" ")[0]
    for rg in pl.REGS:
        wr.cell(r, 1, rg + ("（専任）" if rg in pl.DED else "")).font = F(bold=True); row_border(wr, r)
        wr.cell(r, 2, f'=SUMIFS({LR("F")},{LR("A")},"{rg}")/60').number_format = "0.0"
        for T in pl.open_T:
            if T not in pl.occ[rg] and pl.dem(rg, T) >= 15:
                wr.cell(r, GC + T - G0 // SLOT).fill = PatternFill("solid", fgColor="FFD7D7")
        for T0, T1, nm in blocks(pl.occ[rg]):
            paint(wr, r, T0, T1, short(nm), colp(nm), f"{nm}（{pl.people[nm]['dept']}）{hm(T0*SLOT)}-{hm(T1*SLOT)}")
        r += 1
    for kind in posts:
        lanes_ = []
        for x in sorted([x for x in lrows if x[0] == kind], key=lambda x: x[3]):
            for ln in lanes_:
                if ln[-1][4] <= x[3]: ln.append(x); break
            else: lanes_.append([x])
        for k_, ln in enumerate(lanes_):
            wr.cell(r, 1, kind if k_ == 0 else "").font = F(bold=True); row_border(wr, r)
            if k_ == 0: wr.cell(r, 2, f'=SUMIFS({LR("F")},{LR("A")},"{kind}")/60').number_format = "0.0"
            for _, nm, d, s_, e_ in ln:
                paint(wr, r, s_ // SLOT, e_ // SLOT, short(nm), colp(nm), f"{nm}（{d}）{hm(s_)}-{hm(e_)}")
            r += 1
    r += 1
    wr.cell(r, 1, "凡例").font = F(bold=True)
    wr.cell(r, 2, "薄赤＝予測客数15人/時以上のレジが無人。判定「不足」＝配置台数が必要台数に届かない／「余裕」＝1.5台以上多い").font = F()
    wr.freeze_panes = "D3"

    # ---- タスクリレー割当
    wt.append(["作業指示ID", "作業者氏名", "タイトル", "実施期間F", "実施期間TO", "作業予定時間", f"{md}区分",
               "期間内出勤日数(計画)", "計画外換算日数(週5)", f"{md}割当分", "配置済み分", "持越し分", "備考"])
    for j, (_, t) in enumerate(tr.iterrows(), start=2):
        wt.append([t["作業指示ID"], t["作業者氏名"], t["タイトル"], t["実施期間F"], t["実施期間TO"], t["作業予定時間"],
                   pl.st_day.get(t["nm"], "") or "", t["期間内出勤日数"], t["計画外換算日数"],
                   f'=IF(AND(OR(G{j}="出勤",G{j}="半休"),H{j}+I{j}>0),CEILING(F{j}*60/(H{j}+I{j}),15),0)',
                   f"=SUMIFS(作業配置!$F$2:$F${NA},作業配置!$H$2:$H${NA},A{j})", f"=J{j}-K{j}",
                   "予定時間空欄→中央値" if t["予定時間補完"] else None])
        wt.cell(j, 4).number_format = wt.cell(j, 5).number_format = "yyyy/mm/dd"
        for c in (8, 9): wt.cell(j, c).font = Font(color="0000FF")
    nt = len(tr) + 1
    if len(tr):
        wt.append(["", "合計", "", "", "", f"=SUM(F2:F{nt})", "", "", "", f"=SUM(J2:J{nt})", f"=SUM(K2:K{nt})", f"=SUM(L2:L{nt})"])

    # ---- 作業推定
    we.append(["社員名", "根拠", "根拠の勤務日数", "作業", "実施日数", "1日の中央値(分)", "採用", f"{md}推定分"])
    for row in pl.est_rows: we.append(row)

    # ---- 説明
    S = pl.S
    hist = (f"{pl.hist_from:%Y/%m/%d}〜{pl.hist_to:%Y/%m/%d}" if pl.hist_from is not None else "なし")
    notes = [
        ("項目", "内容"),
        ("店舗", store), ("対象日", f"{pl.day:%Y/%m/%d}({wdn})"), ("行事", pl.event_today or "なし"),
        ("勤務", "月間稼働表（取込済み）の出勤・公休・開始・終了"),
        ("作業の推定", f"KADOU実績 {hist}（直前{S['history_days']}日）。同じ曜日が2日以上あればその曜日、なければ全日。"
                     "最も長い作業＝基本作業、実施率50%以上の作業を中央値の分数で採用。実績の無い人は部門で最も多い作業"),
        ("休憩", "本人の同じ拘束時間の実績→店全体→ルールの順。来客の少ない時間帯に置き、同じ部門で重ならないよう調整"),
        ("来客予測", (f"直前{S['ref_weeks']}週の同じ曜日（{'、'.join(f'{x:%m/%d}' for x in pl.REF)}。祝日・お盆・年末年始・"
                     f"{'/'.join(S['spike_keywords'])}は除外）の時間帯別客数 × 予算売上 ÷ 基準日の平均売上") if pl.has_sales else "売上実績なし"),
        ("有人レジ", "、".join(pl.REGS) + (f"（専任：{'、'.join(pl.DED)}）" if pl.DED else "")),
        ("必要レジ台数", f"有人レジの予測客数 ÷ {S['cap']}人/時 を切上げ、最低{S['min_lanes']}台"),
        ("レジの割当", "専任→レジ部門（予測客数の多いレジから、直前と同じレジは継続）→他部門の応援（台数不足の時間帯へ）"),
        ("1時間交代の持ち場", f"{pl.kero}：営業時間を1時間ずつ1人" if pl.kero else "なし"),
        ("労働時間予算", (f"{pl.labor['getsudo']} {pl.labor['month']:,.1f}h（正社員{pl.labor['regular']:,.1f}・パート{pl.labor['part']:,.1f}・"
                         f"アルバイト{pl.labor['arbeit']:,.1f}）を{pl.labor['method']}で日割り → {pl.labor['day']:.1f}h。"
                         f"計画実働（拘束−休憩）{pl.planned_hours:.1f}h。時間帯予測シートのL8〜M13") if pl.labor else "取込なし"),
        ("タスクリレー", f"実施期間内の出勤予定日数で均等割り（15分切上げ）。1日の上限＝基本作業から{S['tr_keep']}分を残す。"
                        "超過分は持越し。来客の少ない時間帯にまとめて配置"),
    ]
    for x in pl.warnings: notes.append(("要確認", x))
    for x in pl.notes: notes.append(("メモ", x))
    for x in notes: wn.append(x)

    for w in (wf, wt, wl, wa, we, wn):
        for row in w.iter_rows():
            for c in row: c.font = F(size=10, bold=c.row == 1, color=c.font.color)
        for c in w[1]: c.fill = PatternFill("solid", fgColor="DDEBF7")
        w.freeze_panes = "A2"
    for w, ws_ in ((wf, [8, 14, 10, 16, 20, 12, 16, 16, 8, 16, 4, 24, 40]), (wl, [12, 14, 12, 8, 8, 6]),
                   (wa, [12, 14, 12, 8, 8, 6, 30, 9, 50]), (wt, [11, 14, 60, 12, 12, 12, 9, 12, 12, 10, 10, 10, 18]),
                   (we, [14, 26, 12, 12, 9, 14, 10, 10]), (wn, [16, 130])):
        for k, x in enumerate(ws_, 1): w.column_dimensions[get_column_letter(k)].width = x
    for w in (wr, wp):
        w.page_setup.orientation = "landscape"; w.page_setup.paperSize = w.PAPERSIZE_A4
        w.page_setup.fitToWidth, w.page_setup.fitToHeight = 1, 0; w.sheet_properties.pageSetUpPr.fitToPage = True
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    return path
