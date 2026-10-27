"""月度計画（月間稼働表）と労働時間予算の比較・修正案"""
import math
import datetime as dt
import pandas as pd
from .common import WD, m, hm, is_time, is_excluded_day, to_min, merged_settings, getsudo, getsudo_range, getsudo_label

def _rule_break(sh):
    return 0 if sh <= 330 else 60 if sh <= 540 else 90 if sh <= 570 else 120

class BreakModel:
    """休憩：本人の同じ拘束時間の実績 → 店全体の実績 → ルール（日別稼働表と同じ考え方）"""
    def __init__(self, plan, kadou):
        self.own, self.all = {}, {}
        if kadou is None or plan is None or kadou.empty: return
        p = plan[plan["開始"].apply(is_time)].copy()
        p["拘束"] = [m(e) - m(s) for s, e in zip(p["開始"], p["終了"])]
        w = kadou.groupby(["nm", "ymd"])["s_time"].sum().rename("作業").reset_index()
        j = p.merge(w, left_on=["社員名", "日付"], right_on=["nm", "ymd"])
        j["休憩"] = j["拘束"] - j["作業"]
        j = j[j["休憩"] >= 0]
        for (nm, sh), g in j.groupby(["社員名", "拘束"]): self.own[(nm, sh)] = int(g["休憩"].mode().iloc[0]) // 15 * 15
        for sh, g in j.groupby("拘束"): self.all[sh] = int(g["休憩"].mode().iloc[0]) // 15 * 15

    def __call__(self, nm, sh):
        if (nm, sh) in self.own: return self.own[(nm, sh)]
        if sh in self.all: return self.all[sh]
        return _rule_break(sh)

def _demand_profile(data, d0, settings):
    """曜日×時間の来客（直前8週の平均）。売上実績が無ければ一律"""
    hs = data.sales_hour
    o, c = to_min(settings["open"]) // 60, to_min(settings["close"]) // 60
    prof = {}
    if hs is not None and len(hs):
        x = hs[(hs["取引日"] < d0) & (hs["取引日"] >= d0 - pd.Timedelta(days=56))]
        x = x[~x["取引日"].apply(is_excluded_day)]
        if len(x):
            g = x.groupby([x["取引日"].dt.weekday, "h"])["客数"].mean()
            for (wd, h), v in g.items(): prof[(wd, h)] = float(v)
    base = (sorted(prof.values())[len(prof) // 2] if prof else 1.0)
    def dem(wd, h):
        if h < o or h >= c: return base          # 開店前・閉店後は準備・片付けとして中位の扱い
        return prof.get((wd, h), base)
    return dem

def plan_work(pl, brk):
    """月間稼働表の出勤行 → 拘束・休憩・実働（h）"""
    w = pl[pl["区分"].isin(["出勤", "半休"]) & pl["開始"].apply(is_time)].copy()
    w["st"] = w["開始"].apply(m); w["en"] = w["終了"].apply(m)
    w["拘束"] = w["en"] - w["st"]
    w["休憩"] = [brk(n, sh) for n, sh in zip(w["社員名"], w["拘束"])]
    w["実働"] = (w["拘束"] - w["休憩"]) / 60
    return w

def labor_allocation(data, brk):
    """日割りの重み：店の実績（計画時間＝1日あたり固定時間＋売上予算に比例する時間）を当てはめる"""
    bud = data.budget.set_index("日付")["予算売上"] if data.budget is not None else None
    alloc = dict(method="月度の日数で均等割り", a=None, b=None, blend=False, bud=bud)
    if bud is None or data.plan is None: return alloc
    allw = plan_work(data.plan, brk).groupby("日付")["実働"].sum()
    xy = [(float(bud[d]), float(h)) for d, h in allw.items() if d in bud.index and bud[d] > 0]
    if len(xy) >= 20:
        mx = sum(x for x, _ in xy) / len(xy); my = sum(y for _, y in xy) / len(xy)
        vx = sum((x - mx) ** 2 for x, _ in xy)
        b_ = sum((x - mx) * (y - my) for x, y in xy) / vx if vx > 0 else 0
        a_ = my - b_ * mx
        if b_ > 0 and a_ > 0:
            return dict(method=f"店の傾向（1日あたり固定{a_:.0f}h＋売上予算100万円あたり{b_ * 1e6:.1f}h）",
                        a=a_, b=b_, blend=False, bud=bud)
    return dict(method="売上予算の比率（半分）＋日数（半分）", a=None, b=None, blend=True, bud=bud)

def labor_day_budget(data, labor, S, days, brk=None, alloc=None):
    """月度（前月16日〜当月15日）の労働時間予算を日割り。戻り値：({日付: {...}}, 注意書き)"""
    S = merged_settings(S)
    sd = int(S["month_start_day"])
    if brk is None: brk = BreakModel(data.plan, data.kadou)
    if alloc is None: alloc = labor_allocation(data, brk)
    bud = alloc["bud"]
    def weight(d):
        if bud is None or d not in bud.index: return None
        if alloc["a"] is not None: return alloc["a"] + alloc["b"] * float(bud[d])
        if alloc["blend"]: return float(bud[d])
        return None
    days = [pd.Timestamp(d) for d in days]
    out, notes = {}, []
    for ym in sorted({getsudo(d, sd) for d in days}):
        g0, g1 = getsudo_range(ym, sd)
        mdays = list(pd.date_range(g0, g1))
        lb = labor.get(ym)
        if not lb:
            notes.append(f"{getsudo_label(ym, sd)} の労働時間予算がありません"); continue
        ws_ = [weight(x) for x in mdays]
        if all(v is not None for v in ws_) and sum(ws_) > 0:
            tot = sum(ws_)
            if alloc["blend"]: w = {x: 0.5 * v / tot + 0.5 / len(mdays) for x, v in zip(mdays, ws_)}
            else: w = {x: v / tot for x, v in zip(mdays, ws_)}
            how = alloc["method"]
        else:
            w = {x: 1 / len(mdays) for x in mdays}; how = "月度の日数で均等割り"
        for x in mdays:
            if x in days:
                out[x] = dict(total=lb["total"] * w[x], regular=lb["regular"] * w[x], part=lb["part"] * w[x],
                              arbeit=lb["arbeit"] * w[x], how=how, ym=ym, ratio=w[x],
                              month=float(lb["total"]), getsudo=getsudo_label(ym, sd))
    return out, notes

def compare_period(data, labor, settings, d0, d1, tol_rate=0.03, tol_hours=2.0):
    S = merged_settings(settings)
    d0, d1 = pd.Timestamp(d0), pd.Timestamp(d1)
    if data.plan is None: raise ValueError("月間稼働表が取り込まれていません")
    plan = data.plan[(data.plan["日付"] >= d0) & (data.plan["日付"] <= d1)].copy()
    if plan.empty: raise ValueError("指定期間の月間稼働表がありません")
    if "順" in plan: plan = plan.sort_values(["順", "日付"], kind="stable")
    hist = data.kadou[data.kadou["ymd"] < d1 + pd.Timedelta(days=1)] if data.kadou is not None else None
    brk = BreakModel(data.plan, hist)
    dem = _demand_profile(data, d0, S)
    days = list(pd.date_range(d0, d1))
    bud = data.budget.set_index("日付")["予算売上"] if data.budget is not None else None
    events = data.plan.dropna(subset=["行事"]).groupby("日付")["行事"].first() if "行事" in data.plan else {}

    # ---- 計画（人×日）
    work = plan_work(plan, brk)
    # 有休・会議・研修（月間稼働表に時刻が無い日）
    lk = list(S["leave_kinds"]); lh = float(S["leave_hours"] or 0)
    leave = plan[plan["区分"].isin(lk)]
    leave_cnt = leave["区分"].value_counts().reindex(lk).fillna(0).astype(int).to_dict()
    leave_by_day = leave.groupby("日付").size()
    dept_order = list(dict.fromkeys(plan["部門"]))
    person_order = list(dict.fromkeys(plan["社員名"]))
    o_min, c_min = to_min(S["open"]), to_min(S["close"])
    sd = int(S["month_start_day"])

    # ---- 日ごとの労働時間予算（月度の予算を店の傾向で日割り。日別稼働表と同じ計算）
    day_budget, notes = labor_day_budget(data, labor, S, days)

    # ---- 実績（KADOU：作業時間の合計＝休憩を除く実働）
    act_d = {}
    if data.kadou is not None:
        today = pd.Timestamp.today().normalize()      # KADOUは先の日付も予定として入っている → 今日より前だけを実績とする
        k = data.kadou[(data.kadou["ymd"] >= d0) & (data.kadou["ymd"] <= d1) & (data.kadou["ymd"] < today)]
        act_d = (k.groupby("ymd")["s_time"].sum() / 60).to_dict()

    # ---- 修正案
    tol = lambda b: max(tol_hours, b * tol_rate)
    plan_d = work.groupby("日付")["実働"].sum()
    if lh > 0: plan_d = plan_d.add(leave_by_day * lh, fill_value=0)
    gap = {d: float(plan_d.get(d, 0)) - day_budget[d]["total"] for d in days if d in day_budget}
    month_diff = sum(gap.values())
    month_tol = tol(sum(v["total"] for v in day_budget.values()))
    targets = {}
    if not gap:
        mode = "労働時間予算なし"
    elif month_diff > month_tol:
        mode = "削減"           # 月度合計が予算超過 → 超過の大きい日から削る
        need_total = month_diff
        for d, g in sorted(gap.items(), key=lambda x: -x[1]):
            if need_total <= 0 or g <= 0: break
            t = min(g, need_total); targets[d] = t; need_total -= t
        if need_total > 0:      # 日ごとは予算内でも合計が超える → 計画の多い日から均等に
            top = sorted(gap, key=lambda d: -float(plan_d.get(d, 0)))[:10]
            for d in top: targets[d] = targets.get(d, 0) + need_total / len(top)
    else:
        mode = "平準化"         # 月度合計は予算内 → 日の超過分だけ（余裕のある日へ回す）
        targets = {d: g for d, g in gap.items() if g > tol(day_budget[d]["total"])}
    actions = []
    workdays = work.groupby("社員名")["日付"].nunique().to_dict()
    for d, need in sorted(targets.items()):
        w = work[work["日付"] == d]
        cands = []
        for _, r in w.iterrows():
            if r["部門"] == "マネジメント" or r["拘束"] <= 240: continue
            same = w[(w["部門"] == r["部門"]) & (w["社員名"] != r["社員名"])]
            for cut in (30, 60):
                if r["拘束"] - cut < 240: continue
                for side in ("終了", "開始"):
                    # 閉店・開店の担当は、同じ部門で他に1人残る場合だけ
                    if side == "終了" and r["en"] >= c_min and not (same["en"] >= r["en"]).any(): continue
                    if side == "開始" and r["st"] <= o_min and not (same["st"] <= r["st"]).any(): continue
                    st, en = (r["st"], r["en"] - cut) if side == "終了" else (r["st"] + cut, r["en"])
                    nsh = en - st
                    red = (r["拘束"] - r["休憩"] - (nsh - brk(r["社員名"], nsh))) / 60
                    if red <= 0: continue
                    if side == "終了": hrs = list(range(en // 60, math.ceil(r["en"] / 60)))
                    else: hrs = list(range(r["st"] // 60, math.ceil(st / 60)))
                    dm = sum(dem(d.weekday(), h) for h in hrs) / max(1, len(hrs))
                    cands.append(dict(score=dm / red, red=red, nm=r["社員名"], dept=r["部門"], side=side, cut=cut,
                                      hrs=tuple(hrs), now=f"{hm(r['st'])}-{hm(r['en'])}", new=f"{hm(st)}-{hm(en)}", dm=dm))
        used, per_hour = set(), {}
        for c in sorted(cands, key=lambda c: (c["score"], -c["red"])):
            if need <= 0.25: break
            if c["nm"] in used: continue
            if any(per_hour.get((c["side"], h), 0) >= 2 for h in c["hrs"]): continue   # 同じ時間帯を詰めるのは2人まで
            used.add(c["nm"]); need -= c["red"]
            for h in c["hrs"]: per_hour[(c["side"], h)] = per_hour.get((c["side"], h), 0) + 1
            actions.append(dict(日付=d, 種類="短縮" if mode == "削減" else "短縮（平準化）", 部門=c["dept"], 社員名=c["nm"],
                                現在=c["now"], 修正後=c["new"], 増減h=-round(c["red"], 2),
                                理由=f"{c['side']}側の時間帯は来客が少ない（平均{c['dm']:.0f}人/時）"))
        if mode == "削減" and need > 4:   # 短縮だけでは足りない → 公休への変更候補（出勤日数の多い人）
            rest = w[(w["部門"] != "マネジメント") & ~w["社員名"].isin(used)].copy()
            rest["wd"] = rest["社員名"].map(workdays)
            for _, r in rest.sort_values(["wd", "実働"], ascending=False).iterrows():
                if need <= 2: break
                need -= r["実働"]
                actions.append(dict(日付=d, 種類="公休へ変更（候補）", 部門=r["部門"], 社員名=r["社員名"],
                                    現在=f"{hm(r['st'])}-{hm(r['en'])}", 修正後="公休", 増減h=-round(r["実働"], 2),
                                    理由=f"月度の出勤{int(r['wd'])}日と多い"))
    # 余裕のある日：来客の多い時間帯に追加できる時間
    room = []
    for d in days:
        b = day_budget.get(d)
        if not b: continue
        p = float(plan_d.get(d, 0))
        if b["total"] - p > tol(b["total"]):
            top = sorted(range(o_min // 60, c_min // 60), key=lambda h: -dem(d.weekday(), h))[:2]
            room.append(dict(日付=d, 余裕h=round(b["total"] - p, 1), 来客の多い時間帯="・".join(f"{h}時台" for h in sorted(top))))

    # ---- 集計
    act = pd.DataFrame(actions)
    rows = []
    for d in days:
        w = work[work["日付"] == d]
        b = day_budget.get(d)
        cut = float(act.loc[act["日付"] == d, "増減h"].sum()) if len(act) else 0.0
        rows.append(dict(日付=d, 曜日=WD[d.weekday()], 行事=events.get(d) if hasattr(events, "get") else None,
                         売上予算=(float(bud[d]) if bud is not None and d in bud.index else None),
                         出勤人数=int(w["社員名"].nunique()), 有休等=int(leave_by_day.get(d, 0)),
                         計画実働=round(float(w["実働"].sum()) + lh * int(leave_by_day.get(d, 0)), 2),
                         予算=(round(b["total"], 2) if b else None), 修正案=round(cut, 2),
                         実績=(round(float(act_d[d]), 2) if d in act_d else None)))
    daydf = pd.DataFrame(rows)
    dept = work.groupby("部門").agg(人数=("社員名", "nunique"), 出勤延べ=("社員名", "count"), 計画実働=("実働", "sum"))
    dept = dept.reindex([x for x in dept_order if x in dept.index]).reset_index()
    pers = work.groupby(["部門", "社員名"]).agg(出勤日数=("日付", "nunique"), 計画実働=("実働", "sum"),
                                              平均拘束=("拘束", "mean")).reset_index()
    pers["_o"] = pers["社員名"].map({n: i for i, n in enumerate(person_order)})
    pers = pers.sort_values("_o").drop(columns="_o")
    if len(act):
        pc = act.groupby("社員名")["増減h"].sum()
        pers["修正案"] = pers["社員名"].map(pc).fillna(0)
    else:
        pers["修正案"] = 0.0
    btot = sum(v["total"] for v in day_budget.values())
    covered = len(day_budget) == len(days)
    summ = dict(
        period=f"{d0:%Y/%m/%d}〜{d1:%Y/%m/%d}", days=len(days), budget_days=len(day_budget), covered=covered,
        budget=round(btot, 1),
        budget_regular=round(sum(v["regular"] for v in day_budget.values()), 1),
        budget_part=round(sum(v["part"] for v in day_budget.values()), 1),
        budget_arbeit=round(sum(v["arbeit"] for v in day_budget.values()), 1),
        plan=round(float(plan_d.sum()), 1),
        plan_in_budget_days=round(float(plan_d[plan_d.index.isin(list(day_budget))].sum()), 1),
        leave=leave_cnt, leave_hours=lh, leave_days=int(len(leave)),
        actual=round(float(sum(act_d.values())), 1), actual_days=len(act_d),
        actual_budget=round(float(sum(day_budget[d]["total"] for d in act_d if d in day_budget)), 1),
        getsudo=[getsudo_label(ym, sd) for ym in sorted({getsudo(d, sd) for d in days})],
        proposal=round(float(act["増減h"].sum()) if len(act) else 0.0, 1),
        how=sorted({v["how"] for v in day_budget.values()}), notes=notes, mode=mode, month_diff=round(month_diff, 1),
        over_days=int(sum(1 for r in rows if r["予算"] and r["計画実働"] - r["予算"] > tol(r["予算"]))),
        under_days=len(room), room_total=round(sum(x["余裕h"] for x in room), 1),
    )
    summ["diff"] = round(summ["plan_in_budget_days"] - summ["budget"], 1)
    summ["after"] = round(summ["plan_in_budget_days"] + summ["proposal"], 1)
    return dict(summary=summ, days=daydf, dept=dept, person=pers, actions=act, room=pd.DataFrame(room))

# ---------------------------------------------------------------- Excel
def write_excel(res, path, store_name):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    F = lambda **k: Font(name="Meiryo UI", size=k.pop("size", 10), **k)
    S = res["summary"]
    wb = Workbook()
    ws = wb.active; ws.title = "月度比較"
    wd = wb.create_sheet("日別"); wa = wb.create_sheet("修正案"); wp = wb.create_sheet("個人別")
    wdp = wb.create_sheet("部署別"); wr = wb.create_sheet("余裕のある日")

    # 日別（数式の元）
    hdr = ["日付", "曜日", "行事", "売上予算(円)", "出勤人数", "労働時間予算(h)", "計画実働(h)", "差(計画−予算)", "修正案(h)", "修正後(h)", "修正後の差", "有休・会議等(人)", "実績(KADOU・h)", "実績−予算"]
    wd.append(hdr)
    for i, r in enumerate(res["days"].itertuples(), start=2):
        wd.append([r.日付.to_pydatetime(), r.曜日, r.行事, r.売上予算, r.出勤人数, r.予算, r.計画実働,
                   f'=IF(F{i}="","",G{i}-F{i})', f"=SUMIFS(修正案!$H:$H,修正案!$A:$A,A{i})", f"=G{i}+I{i}",
                   f'=IF(F{i}="","",J{i}-F{i})', r.有休等, r.実績, f'=IF(OR(F{i}="",M{i}=""),"",M{i}-F{i})'])
        wd.cell(i, 1).number_format = "yyyy/mm/dd"; wd.cell(i, 4).number_format = "#,##0"
        for c in (6, 7, 8, 9, 10, 11): wd.cell(i, c).number_format = "0.0"
    n = len(res["days"]) + 1
    wd.append(["合計", "", "", f"=SUM(D2:D{n})", f"=SUM(E2:E{n})", f"=SUM(F2:F{n})", f"=SUM(G2:G{n})",
               f"=SUMPRODUCT((F2:F{n}<>\"\")*(G2:G{n}-F2:F{n}))", f"=SUM(I2:I{n})", f"=SUM(J2:J{n})",
               f"=SUMPRODUCT((F2:F{n}<>\"\")*(J2:J{n}-F2:F{n}))"])
    for c in range(4, 12): wd.cell(n + 1, c).number_format = "#,##0.0"
    from openpyxl.formatting.rule import CellIsRule
    wd.conditional_formatting.add(f"H2:H{n}", CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor="FFD7D7")))
    wd.conditional_formatting.add(f"K2:K{n}", CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor="FFD7D7")))

    # 修正案
    wa.append(["日付", "曜日", "種類", "部門", "社員名", "現在", "修正後", "増減(h)", "理由"])
    for i, r in enumerate(res["actions"].itertuples(), start=2) if len(res["actions"]) else []:
        wa.append([r.日付.to_pydatetime(), WD[r.日付.weekday()], r.種類, r.部門, r.社員名, r.現在, r.修正後, r.増減h, r.理由])
        wa.cell(i, 1).number_format = "yyyy/mm/dd"; wa.cell(i, 8).number_format = "0.00"

    # 個人別
    wp.append(["部門", "社員名", "出勤日数", "計画実働(h)", "平均拘束(h)", "修正案(h)", "修正後(h)"])
    for i, r in enumerate(res["person"].itertuples(), start=2):
        wp.append([r.部門, r.社員名, r.出勤日数, round(r.計画実働, 2), round(r.平均拘束 / 60, 2),
                   f"=SUMIFS(修正案!$H:$H,修正案!$E:$E,B{i})", f"=D{i}+F{i}"])
        for c in (4, 5, 6, 7): wp.cell(i, c).number_format = "0.0"
    # 部署別
    wdp.append(["部門", "人数", "出勤延べ", "計画実働(h)", "修正案(h)", "修正後(h)"])
    for i, r in enumerate(res["dept"].itertuples(), start=2):
        wdp.append([r.部門, r.人数, r.出勤延べ, round(r.計画実働, 2), f"=SUMIFS(修正案!$H:$H,修正案!$D:$D,A{i})", f"=D{i}+E{i}"])
        for c in (4, 5, 6): wdp.cell(i, c).number_format = "#,##0.0"
    nd = len(res["dept"]) + 1
    wdp.append(["合計", f"=SUM(B2:B{nd})", f"=SUM(C2:C{nd})", f"=SUM(D2:D{nd})", f"=SUM(E2:E{nd})", f"=SUM(F2:F{nd})"])
    # 余裕のある日
    wr.append(["日付", "曜日", "余裕(h)", "来客の多い時間帯（追加するなら）"])
    for r in res["room"].itertuples() if len(res["room"]) else []:
        wr.append([r.日付.to_pydatetime(), WD[r.日付.weekday()], r.余裕h, r.来客の多い時間帯])
        wr.cell(wr.max_row, 1).number_format = "yyyy/mm/dd"

    # 月度比較（要約）
    L = n + 1
    ws["A1"] = f"月度計画と労働時間予算の比較　{store_name}　{S['period']}　{'・'.join(S['getsudo'])}"; ws["A1"].font = F(size=13, bold=True)
    rows = [
        ("労働時間予算（月度・h）", f"=日別!F{L}"),
        ("　うち正社員", S["budget_regular"]), ("　うちパート", S["budget_part"]), ("　うちアルバイト", S["budget_arbeit"]),
        ("計画実働（h）", f"=日別!G{L}"),
        ("差（計画−予算・h）", f"=日別!H{L}"),
        ("差（%）", f"=IF(B3=0,\"\",B8/B3)"),
        ("修正案の合計（h）", f"=日別!I{L}"),
        ("修正後の計画（h）", f"=日別!J{L}"),
        ("修正後の差（h）", f"=日別!K{L}"),
        ("実績（KADOU・h）", f"=SUM(日別!M2:日別!M{n})"),
        ("実績のある日の予算−実績（h）", f'=SUMPRODUCT((日別!M2:M{n}<>"")*(日別!F2:F{n}))-SUM(日別!M2:M{n})'),
        ("予算超過の日数", S["over_days"]), ("余裕のある日数", S["under_days"]),
        ("日割の方法", "・".join(S["how"]) or "—"),
        ("有休・会議等（日数）", "・".join(f"{k}{v}日" for k, v in S["leave"].items())),
        ("有休・会議等の扱い", f"1日{S['leave_hours']:g}hとして計画に含める" if S["leave_hours"] else "計画に含めない（店舗の設定で変更可）"),
    ]
    for k, (a, b) in enumerate(rows, start=3):
        ws.cell(k, 1, a).font = F(bold=not a.startswith("　")); ws.cell(k, 2, b).font = F()
        ws.cell(k, 2).number_format = "0.0%" if "%" in a else "#,##0.0"
        if isinstance(b, str) and not b.startswith("="): ws.cell(k, 2).alignment = Alignment(horizontal="left")
    r0 = len(rows) + 4
    notes = [
        "計画実働＝月間稼働表の拘束時間−休憩。休憩はKADOU実績から推定（本人の同じ拘束時間→店全体→ルール）",
        "労働時間予算は月度（前月16日〜当月15日。予算の「3月」＝2/16〜3/15）の予算を、店の傾向（固定時間＋売上予算に比例）で日割りしたもの",
        "正社員・パート別の計画は、月間稼働表に身分が無いため比較していません（予算の内訳のみ表示）",
        "修正案：予算超過（3%かつ2h超）の日に、来客の少ない時間帯にかかる勤務の端を30〜60分短縮（マネジメント・4時間以下の勤務は除く）。"
        "足りない日は出勤日数の多い人を公休へ変更する候補。最終判断は店舗で行ってください",
    ] + [f"注意：{x}" for x in S["notes"]]
    for k, x in enumerate(notes): ws.cell(r0 + k, 1, x).font = F(size=9)
    ws.column_dimensions["A"].width = 26; ws.column_dimensions["B"].width = 16
    for w, widths in ((wd, [11, 5, 22, 13, 9, 14, 12, 13, 10, 10, 11, 12, 13, 11]), (wa, [11, 5, 16, 12, 14, 12, 12, 9, 46]),
                      (wp, [12, 14, 9, 12, 12, 10, 10]), (wdp, [12, 8, 9, 12, 10, 10]), (wr, [11, 5, 9, 30])):
        for k, x in enumerate(widths, 1): w.column_dimensions[get_column_letter(k)].width = x
        for row in w.iter_rows():
            for c in row: c.font = F(bold=c.row == 1)
        for c in w[1]: c.fill = PatternFill("solid", fgColor="DDEBF7")
        w.freeze_panes = "A2"
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    return path
